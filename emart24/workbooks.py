from __future__ import annotations

import gc
from contextlib import contextmanager
from pathlib import Path

from .core import StopRun, compact, fingerprint, trim_rows


def read_xlsx(path, sheet=None, source=False):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if book.epoch.year != 1899:
            raise StopRun("1904 날짜 체계 파일은 지원하지 않습니다.")
        ws = book[sheet] if sheet else next(s for s in book if s.sheet_state == "visible")
        headers = [compact(v) for v in next(ws.iter_rows(max_row=1, max_col=15, values_only=True))]
        rows = []
        for row in ws.iter_rows(min_row=2, max_col=15):
            if source and any(c.data_type == "e" for c in row):
                raise StopRun("원본 A~O에 Excel 오류 값이 있습니다.")
            rows.append([c.value for c in row])
        return ws.title, headers, trim_rows(rows)
    finally:
        book.close()


@contextmanager
def excel():
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        app.EnableEvents = False
        app.AskToUpdateLinks = False
        app.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
        yield app
    finally:
        if app is not None:
            app.Quit()
        app = None
        gc.collect()
        pythoncom.CoUninitialize()


def read_source(path, expected):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        title, headers, rows = read_xlsx(path, source=True)
        # Missing cached formula values must not silently become empty cells.
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb[title]
            for row in ws.iter_rows(min_row=2, max_col=15):
                for c in row:
                    if c.data_type == "f" and (c.row - 2 >= len(rows) or rows[c.row - 2][c.column - 1] is None):
                        raise StopRun(f"원본 수식 캐시가 없습니다: {c.coordinate}")
        finally:
            wb.close()
    elif path.suffix.lower() == ".xls":
        with excel() as app:
            wb = app.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=True)
            try:
                if wb.Date1904:
                    raise StopRun("1904 날짜 체계 파일은 지원하지 않습니다.")
                ws = next(s for s in wb.Worksheets if s.Visible == -1)
                title = ws.Name
                headers = [compact(v) for v in ws.Range("A1:O1").Value2[0]]
                end = ws.Range("A:O").Find(What="*", LookIn=-4123, SearchOrder=1, SearchDirection=2)
                if end is None or end.Row < 2:
                    raise StopRun("원본 데이터가 없습니다.")
                rg = ws.Range(f"A2:O{end.Row}")
                if ws.Evaluate(f"SUMPRODUCT(--ISERROR(A2:O{end.Row}))"):
                    raise StopRun("원본 오류 값 확인이 필요합니다.")
                rows = trim_rows(rg.Value2)
            finally:
                wb.Close(False)
    else:
        raise StopRun("지원 형식은 xlsx, xlsm, xls입니다.")
    if headers != [compact(v) for v in expected]:
        raise StopRun("원본 A~O 헤더가 등록된 양식과 다릅니다. 설정을 검토하세요.")
    if not rows:
        raise StopRun("원본에 적재할 데이터가 없습니다.")
    return title, rows


def template_formulas(path, sheet, last):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        values = next(wb[sheet].iter_rows(min_row=last, max_row=last, min_col=16, max_col=20, values_only=True))
        if not all(isinstance(v, str) and v.startswith("=") for v in values):
            raise StopRun("DB 마지막 행의 P~T 수식이 완전하지 않습니다.")
        return values
    finally:
        wb.close()


def append_with_excel(path, sheet, rows, start):
    """Modify only the disposable working copy, never the production file."""
    if start + len(rows) - 1 > 1048576:
        raise StopRun("Excel 최대 행 수를 초과합니다.")
    with excel() as app:
        wb = app.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=False)
        try:
            if wb.ReadOnly:
                raise StopRun("작업 사본이 읽기 전용입니다.")
            app.Calculation = -4135  # manual during writes
            ws = wb.Worksheets(sheet)
            end = start + len(rows) - 1
            ws.Range(f"A{start-1}:T{start-1}").Copy()
            ws.Range(f"A{start}:T{end}").PasteSpecial(Paste=-4122)  # formats only
            app.CutCopyMode = False
            # Text beginning '=' must remain text instead of becoming a formula.
            payload = tuple(tuple("'" + v if isinstance(v, str) else v for v in row) for row in rows)
            ws.Range(f"A{start}:O{end}").Value2 = payload
            for col in range(16, 21):
                formula = ws.Cells(start-1, col).FormulaR1C1
                if not isinstance(formula, str) or not formula.startswith("="):
                    raise StopRun("P~T 수식 확장 기준이 없습니다.")
                ws.Range(ws.Cells(start, col), ws.Cells(end, col)).FormulaR1C1 = formula
            ws.Range(f"P{start}:T{end}").Calculate()
            # Recalculate dependent monthly sheets without refreshing external links.
            app.Calculate()
            wb.Save()
        finally:
            wb.Close(False)


def verify_append(path, sheet, original_rows, added, start, formulas):
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    _, _, result = read_xlsx(path, sheet)
    if len(result) != len(original_rows) + len(added):
        raise StopRun("저장 검증 실패: 행 수가 다릅니다.")
    if fingerprint(result[:start-2]) != fingerprint(original_rows):
        raise StopRun("저장 검증 실패: 기존 데이터가 변경됐습니다.")
    if fingerprint(result[start-2:]) != fingerprint(added):
        raise StopRun("저장 검증 실패: 추가 데이터가 다릅니다.")
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for row in wb[sheet].iter_rows(min_row=start, max_row=start+len(added)-1, min_col=16, max_col=20):
            for c, formula in zip(row, formulas):
                origin = f"{c.column_letter}{start-1}"
                if c.value != Translator(formula, origin=origin).translate_formula(c.coordinate):
                    raise StopRun(f"저장 검증 실패: 수식 참조 {c.coordinate}")
    finally:
        wb.close()
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for row in wb[sheet].iter_rows(min_row=start, max_row=start+len(added)-1, min_col=16, max_col=20):
            if any(c.data_type == "e" or c.value is None for c in row):
                raise StopRun("새 P~T 수식의 계산 결과가 비어 있거나 오류입니다.")
    finally:
        wb.close()
